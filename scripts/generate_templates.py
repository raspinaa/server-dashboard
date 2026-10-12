#!/usr/bin/env python3
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

ROOT = Path(__file__).resolve().parents[1]


def generate():
    wb = Workbook()
    s = wb.active
    s.title = "Inventory"
    headers = [
        "Row",
        "Environment",
        "Application",
        "VMs",
        "CPU (Core)",
        "RAM (GB)",
        "HDD (GB)",
        "HDD (GB)",
        "HDD 3",
        "Disk 4",
        "OS",
        "Network",
        "",
        "Server Name",
        "Server IP Addr",
        "Subnet (CIDR)",
        "GW IP Addr",
        "Username",
        "Password",
    ]
    s.append(headers)
    for i in range(1, 5):
        if i == 3:
            s.append([None] * len(headers))
        disks = [j * 20 if j <= i else None for j in range(1, 5)]
        s.append(
            [
                i,
                "Develop" if i < 3 else "Test",
                "Sample Service",
                "Worker " + str(i),
                2 * i,
                4 * i,
                *disks,
                "Ubuntu 24.04",
                "VLAN 100",
                None,
                f"sample-node-{i}",
                f"192.0.2.{10 + i}",
                "/24",
                "192.0.2.1",
                "sample",
                "",
            ]
        )
    for area in ["B2:B3", "C2:C3", "B5:B6", "C5:C6"]:
        s.merge_cells(area)
    for row in s:
        for c in row:
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = Border(
                left=Side(style="thin", color="CBD5E1"),
                right=Side(style="thin", color="CBD5E1"),
                top=Side(style="thin", color="CBD5E1"),
                bottom=Side(style="thin", color="CBD5E1"),
            )
            if c.row == 1:
                c.font = Font(color="FFFFFF", bold=True)
                c.fill = PatternFill("solid", fgColor="12364A")
    for col in range(1, len(headers) + 1):
        s.column_dimensions[s.cell(1, col).column_letter].width = 22
    s.sheet_view.rightToLeft = True
    s.freeze_panes = "D2"
    ROOT.joinpath("templates").mkdir(exist_ok=True)
    wb.save(ROOT / "templates/infrahub-sanitized.xlsx")
    wb.save(ROOT / "templates/infrahub-template.xlsx")


if __name__ == "__main__":
    generate()
