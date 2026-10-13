import io
import pytest
from openpyxl import Workbook
from rest_framework.exceptions import ValidationError
from core import xlsx
from core.inventory import save_machine, snapshot
from core.models import Machine, Environment
from core.validation import network_fields
from test_xlsx import batch, preview, decision
from test_inventory import data


def grouped_workbook():
    book = Workbook()
    sheet = book.active
    sheet.title = 'Inventory'
    for group, disk_count, ip_count in [('Develop', 1, 1), ('Test', 2, 2), ('Production', 3, 1)]:
        if sheet.max_row > 1:
            sheet.append([])
        headers = ['Row', 'Environment', 'Application', 'VMs', 'CPU (Core)', 'RAM (GB)'] + ['HDD (GB)'] * disk_count + ['Server Name', 'OS', 'Network'] + ['Server IP Addr' + (' ' + str(i + 1) if i else '') for i in range(ip_count)]
        sheet.append(headers)
        start = sheet.max_row + 1
        for i in range(2):
            addresses = [f'192.0.2.{10 + i}/24']
            if ip_count == 2:
                addresses.append(f'2001:db8::{10 + i}/64')
            if group == 'Production':
                addresses[0] += f'\n198.51.100.{10 + i}/25'
            sheet.append([i + 1, group if i == 0 else None, 'Service' if i == 0 else None, f'VM {i}', 2, 4] + [20 + d * 10 for d in range(disk_count)] + [f'node-{i}', 'Ubuntu', 'VLAN'] + addresses)
        sheet.merge_cells(start_row=start, end_row=start + 1, start_column=2, end_column=2)
        sheet.merge_cells(start_row=start, end_row=start + 1, start_column=3, end_column=3)
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def test_variable_sections_multiple_ips_and_roundtrip(setup):
    source = batch(setup, grouped_workbook())
    rows = preview(source)
    assert not source.issues
    assert len(rows) == 6
    assert [r['data']['environment'] for r in rows] == ['Develop'] * 2 + ['Test'] * 2 + ['Production'] * 2
    assert [len(r['data']['disks']) for r in rows] == [1, 1, 2, 2, 3, 3]
    assert [len(r['data']['interfaces']) for r in rows] == [1, 1, 2, 2, 2, 2]
    assert rows[2]['data']['interfaces'][1]['subnet'] == '/64'
    assert rows[4]['data']['interfaces'][1]['subnet'] == '/25'
    xlsx.apply(source, setup[0], decision(rows), source.version)
    assert Machine.objects.count() == 6
    exported = xlsx.export_inventory(setup[1])
    restored = batch(setup, exported)
    parsed = preview(restored, 'add_update')
    assert not restored.issues
    before = {str(m.id): snapshot(m)['interfaces'] for m in Machine.objects.all()}
    assert {row['match']: row['data']['interfaces'] for row in parsed} == before
    xlsx.apply(restored, setup[0], decision(parsed), restored.version)
    assert Machine.objects.count() == 6
    assert {str(m.id): snapshot(m)['interfaces'] for m in Machine.objects.all()} == before


def test_primary_cidr_edit_keeps_secondary_network(setup):
    user, project, env, _ = setup
    machine, _ = save_machine(user, project, data(env, interfaces=[
        {'ip': '192.0.2.10/24', 'network': 'first'},
        {'ip': '198.51.100.10/25', 'network': 'second'},
    ]))
    machine, _ = save_machine(user, project, {'version': machine.version, 'ip': '192.0.2.11/26'}, machine)
    assert machine.ip == '192.0.2.11' and machine.subnet == '/26'
    assert snapshot(machine)['interfaces'][1]['ip'] == '198.51.100.10'


def test_conflicting_explicit_subnet_is_reported():
    with pytest.raises(ValidationError):
        network_fields({'ip': '192.0.2.10/24', 'subnet': '/25'})


def test_single_section_empty_optional_parts_are_zero_or_empty(setup):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Inventory'
    sheet.append(['Environment', 'Server Name', 'CPU (Core)', 'RAM (GB)', 'HDD (GB)', 'Server IP Addr'])
    sheet.append(['Test', 'empty-parts', None, None, None, None])
    output = io.BytesIO()
    book.save(output)
    source = batch(setup, output.getvalue())
    rows = preview(source)
    assert not source.issues and len(rows) == 1
    assert rows[0]['data']['cpu'] == 0
    assert rows[0]['data']['ram'] == '0.00'
    assert rows[0]['data']['disks'] == []
    assert rows[0]['data']['ip'] == ''
    xlsx.apply(source, setup[0], decision(rows), source.version)
    assert Machine.objects.get().cpu == 0


def test_zero_ip_placeholder_is_empty():
    fields = {'ip': '0', 'gateway': '0'}
    assert network_fields(fields) == []
    assert fields['ip'] == fields['gateway'] == ''
