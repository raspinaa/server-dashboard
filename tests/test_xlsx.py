import io, pytest, json
from django.core.files.base import ContentFile
from openpyxl import Workbook, load_workbook
from core import xlsx
from core.models import *
from core.inventory import save_machine, snapshot
from test_inventory import data


def workbook(rows=None, merged=False):
    w = Workbook()
    s = w.active
    s.title = "Inventory"
    s.append(
        [
            "Row",
            "Environment",
            "Application",
            "VMs",
            "CPU (Core)",
            "RAM (GB)",
            "HDD (GB)",
            "HDD (GB)",
            "Disk 3",
            "HDD 4",
            "OS",
            "Network",
            "",
            "Server Name",
            "Server IP Addr",
            "Subnet (CIDR)",
            "GW IP Addr",
            "Username",
            "Password",
        ]
    )
    rows = rows or [
        [
            1,
            "Develop",
            "Demo",
            "Worker 1",
            4,
            8,
            20,
            30,
            0,
            None,
            "Ubuntu 24.04",
            "VLAN20",
            None,
            "node-1",
            "192.0.2.10",
            "24",
            "192.0.2.1",
            "demo",
            "",
        ],
        [
            2,
            "Develop",
            "Demo",
            "Worker 2",
            2,
            4,
            10,
            None,
            None,
            None,
            "Custom OS",
            "VLAN20",
            None,
            "node-2",
            "192.0.2.12",
            "/24",
            "192.0.2.1",
            "",
            "",
        ],
    ]
    for r in rows:
        s.append(r)
    if merged:
        s.merge_cells("B2:B3")
        s.merge_cells("C2:C3")
    s.append([None, "Test"])
    s.append([None] * 19)
    b = io.BytesIO()
    w.save(b)
    return b.getvalue()


def batch(setup, raw):
    a, p, e, _ = setup
    b = ImportBatch.objects.create(project=p, user=a, filename="sample.xlsx")
    b.file.save("sample.encrypted", ContentFile(xlsx.encrypt_file(raw)))
    return b


def preview(b, mode="add_only"):
    info = xlsx.inspect(b)[0]
    return xlsx.preview(b, "Inventory", info["mapping"], info["header_row"], mode)


def decision(rows, **kw):
    return {
        str(r["row"]): {"action": "apply", "fields": "excel", "disks": "replace", **kw}
        for r in rows
    }


def test_merged_repeated_disks_blank_rows(setup):
    b = batch(setup, workbook(merged=True))
    rows = preview(b)
    assert len(rows) == 2
    assert (
        rows[1]["data"]["environment"] == "Develop"
        and rows[1]["data"]["application"] == "Demo"
    )
    assert len(rows[0]["data"]["disks"]) == 2 and len(rows[1]["data"]["disks"]) == 1
    assert not b.issues
    r = xlsx.apply(b, setup[0], decision(rows), b.version)
    assert r == {"added": 2, "updated": 0, "skipped": 0}
    assert Machine.objects.count() == 2 and Disk.objects.count() == 3
    assert (
        ImportDecision.objects.filter(batch=b).count() == 2
        and ImportValidation.objects.filter(batch=b, severity="warning").exists()
    )


@pytest.mark.parametrize(
    "cell,value",
    [
        ("G2", -1),
        ("G2", "20GB"),
        ("E2", 2.5),
        ("O2", "invalid"),
        ("P2", "255.0.255.0"),
        ("E2", "=2+2"),
    ],
)
def test_invalid_import(setup, cell, value):
    wb = load_workbook(io.BytesIO(workbook()))
    wb.active[cell] = value
    s = io.BytesIO()
    wb.save(s)
    b = batch(setup, s.getvalue())
    rows = preview(b)
    assert b.issues
    with pytest.raises(Exception):
        xlsx.apply(b, setup[0], decision(rows), b.version)
    assert not Machine.objects.exists()


def test_matching_requires_decision(setup):
    a, p, e, _ = setup
    save_machine(a, p, data(e, server_name="node-1"))
    save_machine(a, p, data(e, server_name="node-1"))
    b = batch(setup, workbook())
    rows = preview(b, "add_update")
    assert len(rows[0]["candidates"]) == 2
    with pytest.raises(Exception):
        xlsx.apply(b, a, decision(rows), b.version)
    assert Machine.objects.count() == 2


def test_add_only(setup):
    a, p, e, _ = setup
    m, _ = save_machine(a, p, data(e, server_name="node-1"))
    b = batch(setup, workbook())
    rows = preview(b)
    with pytest.raises(Exception):
        xlsx.apply(b, a, decision(rows, match=str(m.id)), b.version)
    decisions = decision(rows)
    decisions["2"] = {"action": "skip"}
    r = xlsx.apply(b, a, decisions, b.version)
    assert r["added"] == 1 and r["skipped"] == 1


def test_conflict_and_disk_preservation(setup):
    a, p, e, _ = setup
    m, _ = save_machine(
        a,
        p,
        data(
            e,
            server_name="node-1",
            cpu=10,
            disks=[
                {"name": "root", "size": 100, "order": 0},
                {"name": "data", "size": 200, "order": 1},
                {"name": "extra", "size": 300, "order": 2},
            ],
        ),
    )
    b = batch(setup, workbook())
    rows = preview(b, "add_update")
    d = decision(rows, match=str(m.id))
    d["2"].update(fields="web", disks="merge")
    xlsx.apply(b, a, d, b.version)
    m.refresh_from_db()
    assert m.cpu == 10
    assert m.disks.count() == 3
    assert m.disks.get(name="extra").size == 300


def test_stale_preview_rollback(setup):
    a, p, e, _ = setup
    m, _ = save_machine(a, p, data(e, server_name="node-2"))
    b = batch(setup, workbook())
    rows = preview(b, "add_update")
    save_machine(a, p, {"version": 1, "cpu": 22}, m)
    d = decision(rows)
    d["3"]["match"] = str(m.id)
    with pytest.raises(Exception):
        xlsx.apply(b, a, d, b.version)
    assert Machine.objects.count() == 1
    assert not Machine.objects.filter(server_name="node-1").exists()
    b.refresh_from_db()
    assert b.status == "preview"


def test_disk_choice_required(setup):
    a, p, e, _ = setup
    m, _ = save_machine(a, p, data(e, server_name="node-1"))
    b = batch(setup, workbook())
    rows = preview(b, "add_update")
    d = decision(rows, match=str(m.id))
    d["2"].pop("disks")
    with pytest.raises(Exception):
        xlsx.apply(b, a, d, b.version)


def test_roundtrip_uuid_and_format(setup):
    a, p, e, _ = setup
    for i in range(1, 5):
        save_machine(
            a,
            p,
            data(
                e,
                server_name=f"node-{i}",
                disks=[
                    {"name": f"volume-{j}", "size": str(j * 10), "order": j - 1}
                    for j in range(1, i + 1)
                ],
            ),
        )
    raw = xlsx.export_inventory(p)
    wb = load_workbook(io.BytesIO(raw))
    ws = wb.active
    info = xlsx.detect(ws)
    assert sum(f == "disk" for f in info["mapping"].values()) == 4
    uuidcol = next(int(c) for c, f in info["mapping"].items() if f == "uuid")
    assert ws.column_dimensions[ws.cell(1, uuidcol).column_letter].hidden
    assert ws.merged_cells.ranges
    b = batch(setup, raw)
    rows = preview(b, "add_update")
    assert all(r["match"] for r in rows)
    assert [len(r["data"]["disks"]) for r in rows] == [1, 2, 3, 4]
    ids = set(Machine.objects.values_list("id", flat=True))
    diskids = set(Disk.objects.values_list("id", flat=True))
    xlsx.apply(b, a, decision(rows), b.version)
    assert ids == set(Machine.objects.values_list("id", flat=True))
    assert diskids == set(Disk.objects.values_list("id", flat=True))
    assert Machine.objects.count() == 4


def test_formula_export_and_activity(setup):
    a, p, e, _ = setup
    m, _ = save_machine(
        a, p, data(e, server_name='=HYPERLINK("bad")', application="@formula")
    )
    Activity.objects.create(machine=m, author=a, event="comment", comment="=danger")
    for raw in [xlsx.export_inventory(p), xlsx.export_activity(p)]:
        wb = load_workbook(io.BytesIO(raw))
        assert not any(c.data_type == "f" for s in wb for row in s for c in row)
    assert xlsx.safe_cell("=bad") == "'=bad"


def test_import_api_redacts_secret(setup, client):
    wb = load_workbook(io.BytesIO(workbook()))
    wb.active["S2"] = "SYNTHETIC-import-secret"
    s = io.BytesIO()
    wb.save(s)
    b = batch(setup, s.getvalue())
    preview(b)
    r = client.get(f"/api/imports/{b.id}/")
    assert r.status_code == 200
    assert "SYNTHETIC-import-secret" not in str(r.data)
    assert "secret" not in r.data["rows"][0]
    with b.file.open("rb") as f:
        assert "SYNTHETIC-import-secret" not in f.read().decode(errors="ignore")


def test_export_revocation(setup):
    from rest_framework.test import APIClient

    a, p, e, _ = setup
    save_machine(a, p, data(e))
    u = User.objects.create_user("exporter")
    member = Membership.objects.create(
        user=u, project=p, role="operator", permissions=["export"]
    )
    c = APIClient()
    c.force_authenticate(u)
    assert (
        c.post(f"/api/projects/{p.id}/exports/", {}, format="json").status_code == 200
    )
    member.permissions = []
    member.save()
    assert (
        c.post(f"/api/projects/{p.id}/exports/", {}, format="json").status_code == 403
    )


def test_preview_version(setup):
    b = batch(setup, workbook())
    rows = preview(b)
    v = b.version
    preview(b)
    with pytest.raises(Exception):
        xlsx.apply(b, setup[0], decision(rows), v)


@pytest.mark.parametrize("name", ["=formula", "'=literal", "+text", "@name"])
def test_formula_text_roundtrip(setup, name):
    a, p, e, _ = setup
    m, _ = save_machine(a, p, data(e, server_name=name))
    b = batch(setup, xlsx.export_inventory(p))
    rows = preview(b, "add_update")
    assert rows[0]["data"]["server_name"] == name
    assert not b.issues


def test_original_template_subnet_extension_position(setup):
    a, p, e, _ = setup
    wb = load_workbook(io.BytesIO(workbook()))
    wb.active.delete_cols(16)
    s = io.BytesIO()
    wb.save(s)
    b = batch(setup, s.getvalue())
    rows = preview(b)
    xlsx.apply(b, a, decision(rows), b.version)
    out = load_workbook(io.BytesIO(xlsx.export_inventory(p)))
    labels = [c.value for c in out.active[1]]
    assert labels.index("Server IP Addr") + 1 == labels.index("Subnet (CIDR)")
    assert labels.index("Subnet (CIDR)") + 1 == labels.index("GW IP Addr")


def test_duplicate_ip_context_is_warning(setup):
    wb = load_workbook(io.BytesIO(workbook()))
    wb.active["O3"] = "192.0.2.10"
    s = io.BytesIO()
    wb.save(s)
    b = batch(setup, s.getvalue())
    rows = preview(b)
    assert any("IP تکراری" in x for x in rows[1]["warnings"]) and not b.issues
